"""
仓库管理视图
"""
import logging
import io
from decimal import Decimal, InvalidOperation
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    TransferOrder, DisposalPlan, LegalFreeze, FreezeItem, FreezeEvent,
    get_blocking_items,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    TransferOrderSerializer, DisposalPlanSerializer,
    LegalFreezeSerializer, FreezeEventSerializer,
    LegalFreezeCreateSerializer, LegalFreezeLiftSerializer,
)
from . import services
from .services import OPERATION_LABELS, OP_STOCK_OUT, OP_TRANSFER, OP_DESTROY
from apps.core.exceptions import BusinessException

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


# ==================== 货物 / 入库 / 预警 / 审批查询 ====================

class GoodsListView(APIView):
    """货物列表视图（含当前冻结阻断依据）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Goods.objects.select_related(
            'variety__category__unit'
        ).filter(is_active=True).order_by('-created_at')

        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        goods_list = list(queryset[start:end])

        # 当前阻断依据：任一生效冻结限制任一受控操作即视为被冻结
        blockers = {}
        for op in (OP_STOCK_OUT, OP_TRANSFER, OP_DESTROY):
            for goods_id, items in get_blocking_items(
                [g.pk for g in goods_list], op
            ).items():
                blockers.setdefault(goods_id, []).extend(items)

        data = GoodsSerializer(goods_list, many=True).data
        for row, goods in zip(data, goods_list):
            items = blockers.get(goods.pk, [])
            row['is_frozen'] = bool(items)
            # 同一冻结可能限制多个操作，按冻结编号去重展示
            seen, basis = set(), []
            for item in items:
                if item.freeze_id in seen:
                    continue
                seen.add(item.freeze_id)
                basis.append({
                    'freeze_id': item.freeze_id,
                    'freeze_no': item.freeze.freeze_no,
                    'case_info': item.freeze.case_info,
                    'authority': item.freeze.authority,
                    'legal_doc': item.freeze.legal_doc,
                    'effective_from': item.freeze.effective_from,
                })
            row['freeze_basis'] = basis

        return success_response(data={
            'list': data, 'total': total, 'page': page, 'page_size': page_size
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related('goods', 'operator').order_by('-stock_in_time')
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': StockInSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Warning.objects.select_related('goods').order_by('-created_at')
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': WarningSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Approval.objects.select_related(
            'stock_out__goods', 'approver'
        ).order_by('-created_at')
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': ApprovalSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })


# ==================== 出库（领用）流程 ====================

class StockOutListView(APIView):
    """出库申请列表 / 提交领用申请"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockOut.objects.select_related('goods', 'operator').order_by('-created_at')
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': StockOutSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        goods_id = request.data.get('goods')
        try:
            goods = Goods.objects.get(pk=goods_id)
        except (TypeError, ValueError, Goods.DoesNotExist):
            return error_response(message='物资不存在', code=404)

        receiver = (request.data.get('receiver') or '').strip()
        if not receiver:
            return error_response(message='请填写领用人')
        try:
            quantity = Decimal(request.data.get('quantity'))
        except (TypeError, ValueError, InvalidOperation):
            return error_response(message='出库数量格式错误')

        try:
            stock_out = services.create_stock_out(
                user=request.user, goods=goods, quantity=quantity,
                receiver=receiver,
                receiver_dept=request.data.get('receiver_dept', ''),
                remark=request.data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)

        logger.info("User %s submitted stock out %s", request.user.username, stock_out.id)
        return success_response(data=StockOutSerializer(stock_out).data, message='申请已提交')


class StockOutReviewView(APIView):
    """出库审批（放行点：审批通过前持锁复查冻结）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        approved = bool(request.data.get('approved'))
        remark = request.data.get('remark', '')
        try:
            stock_out, _ = services.review_stock_out(
                pk, approver=request.user, approved=approved, remark=remark
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        action = '通过' if approved else '拒绝'
        logger.info("User %s %s stock out %s", request.user.username, action, pk)
        return success_response(data=StockOutSerializer(stock_out).data, message=f'已{action}')


class StockOutCompleteView(APIView):
    """执行出库（放行点：扣减库存前再次持锁复查冻结）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            stock_out = services.complete_stock_out(pk, user=request.user)
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        logger.info("User %s completed stock out %s", request.user.username, pk)
        return success_response(data=StockOutSerializer(stock_out).data, message='出库完成')


# ==================== 转移流程 ====================

class TransferListCreateView(APIView):
    """转移记录列表 / 提交转移申请"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = TransferOrder.objects.select_related('goods', 'operator').order_by('-created_at')
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': TransferOrderSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        try:
            goods = Goods.objects.get(pk=request.data.get('goods'))
        except (TypeError, ValueError, Goods.DoesNotExist):
            return error_response(message='物资不存在', code=404)
        target_location = (request.data.get('target_location') or '').strip()
        if not target_location:
            return error_response(message='请填写目标存放位置')
        try:
            quantity = Decimal(request.data.get('quantity'))
        except (TypeError, ValueError, InvalidOperation):
            return error_response(message='转移数量格式错误')
        try:
            order = services.create_transfer(
                user=request.user, goods=goods, quantity=quantity,
                target_location=target_location,
                target_keeper=request.data.get('target_keeper', ''),
                remark=request.data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=TransferOrderSerializer(order).data, message='申请已提交')


class TransferReviewView(APIView):
    """转移审批"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            order = services.review_transfer(
                pk, approver=request.user,
                approved=bool(request.data.get('approved')),
                remark=request.data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=TransferOrderSerializer(order).data)


class TransferCompleteView(APIView):
    """执行转移"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            order = services.complete_transfer(pk, user=request.user)
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=TransferOrderSerializer(order).data, message='转移完成')


# ==================== 销毁流程 ====================

class DisposalListCreateView(APIView):
    """销毁计划列表 / 制定销毁计划"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = DisposalPlan.objects.select_related('goods', 'operator').order_by('-created_at')
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': DisposalPlanSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        try:
            goods = Goods.objects.get(pk=request.data.get('goods'))
        except (TypeError, ValueError, Goods.DoesNotExist):
            return error_response(message='物资不存在', code=404)
        reason = (request.data.get('reason') or '').strip()
        if not reason:
            return error_response(message='请填写销毁事由')
        try:
            quantity = Decimal(request.data.get('quantity'))
        except (TypeError, ValueError, InvalidOperation):
            return error_response(message='销毁数量格式错误')
        planned_raw = request.data.get('planned_time')
        planned_time = None
        if planned_raw:
            from django.utils.dateparse import parse_datetime
            planned_time = parse_datetime(str(planned_raw))
            if planned_time is None:
                return error_response(message='计划销毁时间格式错误')
        try:
            plan = services.create_disposal(
                user=request.user, goods=goods, quantity=quantity,
                reason=reason, planned_time=planned_time,
                remark=request.data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=DisposalPlanSerializer(plan).data, message='销毁计划已提交')


class DisposalReviewView(APIView):
    """销毁审批"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            plan = services.review_disposal(
                pk, approver=request.user,
                approved=bool(request.data.get('approved')),
                remark=request.data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=DisposalPlanSerializer(plan).data)


class DisposalDestroyView(APIView):
    """执行销毁"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            plan = services.destroy_disposal(pk, user=request.user)
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=DisposalPlanSerializer(plan).data, message='销毁完成')


# ==================== 法律冻结 ====================

def _first_serializer_error(errors):
    """提取序列化器首个错误文案"""
    first = list(errors.values())[0]
    if isinstance(first, dict):
        return _first_serializer_error(first)
    if isinstance(first, list):
        first = first[0]
        if isinstance(first, dict):
            return _first_serializer_error(first)
    return str(first)


class FreezeListCreateView(APIView):
    """冻结记录列表 / 建立法律冻结"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = LegalFreeze.objects.prefetch_related(
            'items__goods'
        ).order_by('-effective_from', '-created_at')

        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        case_info = request.query_params.get('case_info')
        if case_info:
            queryset = queryset.filter(case_info__contains=case_info)
        freeze_no = request.query_params.get('freeze_no')
        if freeze_no:
            queryset = queryset.filter(freeze_no__contains=freeze_no)
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(items__goods_id=goods_id).distinct()

        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
        total = queryset.count()
        records = queryset[(page - 1) * page_size: (page - 1) * page_size + page_size]
        return success_response(data={
            'list': LegalFreezeSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        serializer = LegalFreezeCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_serializer_error(serializer.errors))
        data = serializer.validated_data
        try:
            freeze = services.create_freeze(
                user=request.user,
                freeze_no=data['freeze_no'],
                case_info=data['case_info'],
                authority=data['authority'],
                legal_doc=data.get('legal_doc', ''),
                action_type=data.get('action_type') or LegalFreeze.ACTION_FREEZE,
                effective_from=data['effective_from'],
                effective_to=data.get('effective_to'),
                items=[
                    {
                        'goods_id': item['goods'],
                        'operations': item.get('operations') or FreezeItem.ALL_OPERATIONS,
                        'quantity': item.get('quantity'),
                    }
                    for item in data['items']
                ],
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(
            data=LegalFreezeSerializer(freeze).data, message='冻结已建立'
        )


class FreezeDetailView(APIView):
    """冻结详情：当前阻断依据 + 完整时间线"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        freeze = LegalFreeze.objects.prefetch_related(
            'items__goods', 'events__operator', 'events__goods'
        ).filter(pk=pk).first()
        if freeze is None:
            return error_response(message='冻结记录不存在', code=404)

        data = LegalFreezeSerializer(freeze).data
        # 逐份独立计算当前阻断依据
        now_basis = []
        for item in freeze.items.all():
            if freeze.is_effective and item.restrict_operations:
                now_basis.append({
                    'goods_id': item.goods_id,
                    'goods_name': item.goods.name,
                    'goods_code': item.goods.code,
                    'operations': item.operations_list,
                    'operations_display': [
                        OPERATION_LABELS.get(op, op) for op in item.operations_list
                    ],
                    'quantity': str(item.quantity) if item.quantity is not None else None,
                })
        data['current_restrictions'] = now_basis
        data['timeline'] = FreezeEventSerializer(
            freeze.events.all().order_by('occurred_at', 'id'), many=True
        ).data
        return success_response(data=data)


class FreezeLiftView(APIView):
    """解除一份冻结（不影响其他重叠冻结）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = LegalFreezeLiftSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_serializer_error(serializer.errors))
        data = serializer.validated_data
        try:
            freeze = services.lift_freeze(
                pk, user=request.user,
                lift_doc=data.get('lift_doc', ''),
                lift_reason=data.get('lift_reason', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        freeze = LegalFreeze.objects.prefetch_related('items__goods').get(pk=freeze.pk)
        return success_response(data=LegalFreezeSerializer(freeze).data, message='冻结已解除')


class GoodsFreezeStatusView(APIView):
    """物资冻结状态查询：承办人查看当前阻断依据与该物资的冻结时间线"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            goods = Goods.objects.get(pk=pk)
        except Goods.DoesNotExist:
            return error_response(message='物资不存在', code=404)

        by_operation = {}
        all_items = {}
        for op, label in (
            (OP_STOCK_OUT, '出库'), (OP_TRANSFER, '转移'), (OP_DESTROY, '销毁')
        ):
            items = get_blocking_items([goods.pk], op).get(goods.pk, [])
            by_operation[op] = {
                'label': label,
                'blocked': bool(items),
                'basis': [
                    {
                        'freeze_id': item.freeze_id,
                        'freeze_no': item.freeze.freeze_no,
                        'case_info': item.freeze.case_info,
                        'authority': item.freeze.authority,
                        'legal_doc': item.freeze.legal_doc,
                        'action_type': item.freeze.action_type,
                        'effective_from': item.freeze.effective_from,
                        'effective_to': item.freeze.effective_to,
                        'quantity': str(item.quantity) if item.quantity is not None else None,
                    }
                    for item in items
                ],
            }
            for item in items:
                all_items[item.freeze_id] = item.freeze

        # 该物资相关的全部冻结（含已解除、已过期）与事件时间线
        freeze_records = LegalFreeze.objects.filter(
            items__goods=goods
        ).prefetch_related('items').distinct().order_by('-effective_from')
        freezes_data = []
        for freeze in freeze_records:
            item = next(i for i in freeze.items.all() if i.goods_id == goods.pk)
            freezes_data.append({
                'freeze_id': freeze.pk,
                'freeze_no': freeze.freeze_no,
                'case_info': freeze.case_info,
                'authority': freeze.authority,
                'legal_doc': freeze.legal_doc,
                'status': freeze.status,
                'status_display': freeze.get_status_display(),
                'is_effective': freeze.is_effective,
                'effective_from': freeze.effective_from,
                'effective_to': freeze.effective_to,
                'lifted_at': freeze.lifted_at,
                'lift_reason': freeze.lift_reason,
                'operations': item.operations_list,
            })

        events = FreezeEvent.objects.filter(
            goods=goods
        ).select_related('freeze', 'operator').order_by('occurred_at', 'id')

        return success_response(data={
            'goods': GoodsSerializer(goods).data,
            'is_frozen': bool(all_items),
            'restrictions': by_operation,
            'freezes': freezes_data,
            'timeline': FreezeEventSerializer(events, many=True).data,
        })
